import io
import json
import re
import sys
import threading
import urllib.error
import urllib.request
from http.server import HTTPServer
import unittest
import zipfile
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
import cadence_pii_guard as guard


class SecurityRegressionTests(unittest.TestCase):
    def test_generic_account_prose_is_not_redacted_as_bank_account(self):
        text = "The account becomes linked to the risk when triggered."
        _, counts, _ = guard.redact_text(text)
        self.assertNotIn("BANK_ACCOUNT", counts)

    def test_explicit_account_number_is_redacted(self):
        redacted, counts, _ = guard.redact_text("Account number: 123456789012")
        self.assertEqual(counts["BANK_ACCOUNT"], 1)
        self.assertNotIn("123456789012", redacted)

    def test_government_id_header_without_value_is_not_pii(self):
        _, counts, _ = guard.redact_text("Passport number | Passport number: P12345678")
        self.assertEqual(counts["GOVERNMENT_ID"], 1)

    def test_email_with_sentence_period_is_detected_without_punctuation(self):
        text = "Contact riley.example@example.test."
        _, counts, locations = guard.redact_text(text)
        self.assertEqual(counts["EMAIL"], 1)
        self.assertEqual(text[locations[0]["start"]:locations[0]["end"]], "riley.example@example.test")

    def test_unlabeled_digits_are_not_assumed_to_be_phone_numbers(self):
        _, counts, _ = guard.redact_text("Policy reference sequence 12345678901")
        self.assertNotIn("PHONE", counts)
        _, labeled_counts, _ = guard.redact_text("Mobile: 12345678901")
        self.assertEqual(labeled_counts["PHONE"], 1)

    def test_csv_headers_are_attached_to_contact_values(self):
        text = guard.extract_csv(b"Full Name,Email,Mobile\nRiley Example,riley.example@example.test,+1 (202) 555-0199\n")
        _, counts, _ = guard.redact_text(text)
        self.assertEqual(counts["PERSON_NAME"], 1)
        self.assertEqual(counts["EMAIL"], 1)
        self.assertEqual(counts["PHONE"], 1)

    def test_csv_extra_cells_raise_and_can_be_quarantined(self):
        with self.assertRaises(ValueError):
            guard.extract_document("bad.csv", b"Name,Email\nRiley Example,riley@example.test,EXTRA", False)

    def test_resume_header_offsets_preserve_location_marker(self):
        text = "[PARAGRAPH 1] Riley Sample\n[PARAGRAPH 2] person@example.test"
        redacted, counts, _ = guard.redact_text(text)
        self.assertEqual(counts["PERSON_NAME"], 1)
        self.assertIn("[PARAGRAPH 1] [REDACTED:PERSON_NAME]", redacted)
        self.assertNotIn("Riley Sample", redacted)

    def test_table_header_context_catches_short_demo_mobile_values(self):
        text = "[SLIDE 8 TABLE ROW 2] Name: Riley Sample | Email: person@example.test | Mobile: +1 (555) 0114"
        redacted, counts, _ = guard.redact_text(text)
        self.assertGreaterEqual(counts["PHONE"], 1)
        self.assertNotIn("555) 0114", redacted)

    def test_docx_table_headers_remain_attached_to_values(self):
        from docx import Document
        stream = io.BytesIO()
        doc = Document()
        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "Name"
        table.cell(0, 1).text = "Mobile"
        table.cell(1, 0).text = "Riley Sample"
        table.cell(1, 1).text = "+1 (555) 0114"
        doc.save(stream)
        text = guard.extract_docx(stream.getvalue())
        redacted, counts, _ = guard.redact_text(text)
        self.assertIn("Mobile: +1 (555) 0114", text)
        self.assertGreaterEqual(counts["PHONE"], 1)
        self.assertNotIn("555) 0114", redacted)

    def test_pdf_extension_spoof_is_quarantined(self):
        _, reasons = guard.extract_document("fake.pdf", b"not a pdf", False)
        self.assertTrue(reasons)

    def test_real_docx_renamed_pdf_and_real_pdf_renamed_docx_are_quarantined(self):
        from docx import Document
        from pypdf import PdfWriter
        docx_stream = io.BytesIO()
        Document().save(docx_stream)
        _, docx_as_pdf_reasons = guard.extract_document("renamed.pdf", docx_stream.getvalue(), False)
        self.assertTrue(docx_as_pdf_reasons)
        pdf_stream = io.BytesIO()
        writer = PdfWriter(); writer.add_blank_page(width=72, height=72); writer.write(pdf_stream)
        _, pdf_as_docx_reasons = guard.extract_document("renamed.docx", pdf_stream.getvalue(), False)
        self.assertTrue(pdf_as_docx_reasons)

    def test_empty_file_is_quarantined(self):
        _, reasons = guard.extract_document("empty.pdf", b"", False)
        self.assertTrue(reasons)

    def test_invalid_utf8_is_quarantined_without_replacement(self):
        text, reasons = guard.extract_document("bad.txt", b"name=\xff", False)
        self.assertEqual(text, "")
        self.assertTrue(reasons)

    def test_malformed_office_container_is_quarantined(self):
        valid, reasons = guard.inspect_office_package(b"not a zip", ".docx")
        self.assertFalse(valid)
        self.assertTrue(reasons)

    def test_high_ratio_office_zip_is_rejected(self):
        stream = io.BytesIO()
        with zipfile.ZipFile(stream, "w", compression=zipfile.ZIP_DEFLATED) as archive:
            archive.writestr("word/document.xml", "A" * 1_000_000)
        valid, reasons = guard.inspect_office_package(stream.getvalue(), ".docx")
        self.assertFalse(valid)
        self.assertIn("decompression", reasons[0])

    def test_supplied_pptx_mobile_column_is_preserved_and_redacted(self):
        source = ROOT / "Sample Files" / "Cadence_Financial_Group_Organizational_Pack.pptx"
        if not source.exists():
            self.skipTest("Sample pack not present")
        text = guard.extract_pptx(source.read_bytes())
        self.assertIn("Mobile: +1 (555) 0114", text)
        redacted, counts, _ = guard.redact_text(text)
        self.assertGreaterEqual(counts["PHONE"], 16)
        self.assertNotIn("555) 0114", redacted)
        self.assertNotIn("BANK_ACCOUNT", counts)

    def test_local_upload_requires_same_origin_and_csrf_token(self):
        server = HTTPServer(("127.0.0.1", 0), guard.Handler)
        thread = threading.Thread(target=server.serve_forever, daemon=True)
        thread.start()
        base = f"http://127.0.0.1:{server.server_address[1]}"
        try:
            with urllib.request.urlopen(base + "/") as response:
                page = response.read().decode()
                self.assertEqual(response.headers.get("Cache-Control"), "no-store")
                token = re.search(r"X-CSRF-Token':'([^']+)'", page).group(1)
            boundary = "codex-test-boundary"
            body = (f"--{boundary}\r\nContent-Disposition: form-data; name=\"files\"; filename=\"sample.txt\"\r\n"
                    "Content-Type: text/plain\r\n\r\ncontact: person@example.test\r\n"
                    f"--{boundary}--\r\n").encode()
            request = urllib.request.Request(base + "/analyze", data=body, method="POST", headers={
                "Content-Type": f"multipart/form-data; boundary={boundary}",
                "Origin": base,
            })
            with self.assertRaises(urllib.error.HTTPError) as blocked:
                urllib.request.urlopen(request)
            self.assertEqual(blocked.exception.code, 403)
            request.add_header("X-CSRF-Token", token)
            with urllib.request.urlopen(request) as response:
                result = json.loads(response.read())
                self.assertEqual(result["files"][0]["counts"]["EMAIL"], 1)
                self.assertNotIn("person@example.test", result["files"][0]["redacted"])
        finally:
            server.shutdown()
            server.server_close()
            thread.join(timeout=2)


if __name__ == "__main__":
    unittest.main()
